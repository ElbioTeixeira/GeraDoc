import os
import re
import copy
import threading
import pandas as pd
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

# Importa win32com para automação nativa e segura do Word
try:
    import win32com.client
    import pythoncom
    HAS_WIN32COM = True
except ImportError:
    HAS_WIN32COM = False

def higienizar_nome_arquivo(nome):
    if not nome or str(nome).strip() == 'nan':
        return "documento_sem_nome"
    nome_limpo = re.sub(r'[\\/*?":"<>|]', '_', str(nome).strip())
    return nome_limpo

# --- Função de Conversão Direta para PDF via COM (Sem depender de Terminal) ---
def converter_para_pdf_seguro(caminho_docx):
    """
    Abre o Word em segundo plano via COM do Windows, converte para PDF e fecha,
    sem tentar escrever logs em terminais inexistentes.
    """
    if not HAS_WIN32COM:
        raise ImportError("A biblioteca pywin32 não está instalada.")
    
    caminho_docx = os.path.abspath(caminho_docx)
    caminho_pdf = os.path.splitext(caminho_docx)[0] + ".pdf"
    
    pythoncom.CoInitialize()
    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0 # Ignora alertas para não travar a tela
    
    try:
        doc = word.Documents.Open(caminho_docx)
        # Formato 17 corresponde ao PDF no Word
        doc.SaveAs2(caminho_pdf, FileFormat=17)
        doc.Close()
    finally:
        word.Quit()
        pythoncom.CoUninitialize()

# --- Funções de Formatação (Máscaras de Dados) ---

def aplicar_mascara_cpf(valor):
    apenas_numeros = re.sub(r'\D', '', str(valor))
    if len(apenas_numeros) == 11:
        return f"{apenas_numeros[:3]}.{apenas_numeros[3:6]}.{apenas_numeros[6:9]}-{apenas_numeros[9:]}"
    return str(valor)

def aplicar_mascara_cnpj(valor):
    apenas_numeros = re.sub(r'\D', '', str(valor))
    if len(apenas_numeros) == 14:
        return f"{apenas_numeros[:2]}.{apenas_numeros[2:5]}.{apenas_numeros[5:8]}/{apenas_numeros[8:12]}-{apenas_numeros[12:]}"
    return str(valor)

def aplicar_mascara_telefone(valor):
    apenas_numeros = re.sub(r'\D', '', str(valor))
    if len(apenas_numeros) == 11:
        return f"({apenas_numeros[:2]}) {apenas_numeros[2:7]}-{apenas_numeros[7:]}"
    elif len(apenas_numeros) == 10:
        return f"({apenas_numeros[:2]}) {apenas_numeros[2:6]}-{apenas_numeros[6:]}"
    return str(valor)

def aplicar_mascara_moeda(valor):
    try:
        num = float(str(valor).replace(',', '.').strip())
        return f"R$ {num:,.2f}".replace(',', 'v').replace('.', ',').replace('v', '.')
    except ValueError:
        return str(valor)

def processar_regras_formatacao(nome_coluna, regra, row_dados):
    valor_bruto = str(row_dados.get(nome_coluna, '')).strip() if pd.notna(row_dados.get(nome_coluna, '')) else ""

    if not regra:
        return valor_bruto

    regra_limpa = regra.lower().strip()
    if regra_limpa == 'cpf':
        return aplicar_mascara_cpf(valor_bruto)
    elif regra_limpa == 'cnpj':
        return aplicar_mascara_cnpj(valor_bruto)
    elif regra_limpa == 'fone' or regra_limpa == 'telefone':
        return aplicar_mascara_telefone(valor_bruto)
    elif regra_limpa == 'moeda' or regra_limpa == 'dinheiro':
        return aplicar_mascara_moeda(valor_bruto)

    return valor_bruto

def extrair_e_substituir_tags(texto_original, row_dados):
    if "{{" not in texto_original:
        return texto_original

    texto_final = texto_original
    tags_encontradas = re.findall(r'\{\{([^}]+)\}\}', texto_original)

    for tag_interna in tags_encontradas:
        regra = None
        coluna_pura = tag_interna

        if ":" in tag_interna:
            partes = tag_interna.split(':', 1)
            coluna_pura = partes[0].strip()
            regra = partes[1].strip()

        if coluna_pura in row_dados:
            valor_final = processar_regras_formatacao(coluna_pura, regra, row_dados)
            texto_final = texto_final.replace(f"{{{{{tag_interna}}}}}", valor_final)

    return texto_final

def substituir_preservando_formatacao(paragrafo, row_dados):
    texto_original = paragrafo.text
    if "{{" not in texto_original:
        return

    texto_formatado = extrair_e_substituir_tags(texto_original, row_dados)

    for run in paragrafo.runs:
        if "{{" in run.text:
            run.text = extrair_e_substituir_tags(run.text, row_dados)

    if "{{" in paragrafo.text:
        if paragrafo.runs:
            primeiro_run = paragrafo.runs[0]
            primeiro_run.text = texto_formatado
            for run in paragrafo.runs[1:]:
                run.text = ""
        else:
            paragrafo.text = texto_formatado

def duplicar_paragrafo_estilizado(paragrafo_modelo, paragrafo_ref):
    novo_p_element = OxmlElement('w:p')
    paragrafo_ref._p.getparent().insert(paragrafo_ref._p.getparent().index(paragrafo_ref._p) + 1, novo_p_element)

    novo_p = Paragraph(novo_p_element, paragrafo_ref._parent)

    if paragrafo_modelo._p.pPr is not None:
        novo_p._p.append(OxmlElement('w:pPr'))
        for child in paragrafo_modelo._p.pPr:
            novo_p._p.pPr.append(child.__copy__())

    for run_modelo in paragrafo_modelo.runs:
        novo_run = novo_p.add_run()
        novo_run.text = run_modelo.text
        if run_modelo._r.rPr is not None:
            novo_run._r.append(OxmlElement('w:rPr'))
            for child in run_modelo._r.rPr:
                novo_run._r.rPr.append(child.__copy__())

    return novo_p

def inserir_paragrafo_depois(paragrafo_ref, texto=""):
    novo_p_element = OxmlElement('w:p')
    paragrafo_ref._p.getparent().insert(paragrafo_ref._p.getparent().index(paragrafo_ref._p) + 1, novo_p_element)
    novo_p = Paragraph(novo_p_element, paragrafo_ref._parent)
    novo_p.text = texto
    return novo_p

# --- Funções de Junção de Arquivos ---
def combinar_documentos(arquivos_gerados, caminho_saida):
    if not arquivos_gerados:
        return

    doc_final = Document(arquivos_gerados[0])

    for arq in arquivos_gerados[1:]:
        doc_final.add_page_break()
        doc_temp = Document(arq)
        
        sectPr_final = doc_final._body._element.find(qn('w:sectPr'))
        
        for element in doc_temp._body._element:
            if element.tag.endswith('sectPr'):
                continue
            
            if sectPr_final is not None:
                sectPr_final.addprevious(copy.deepcopy(element))
            else:
                doc_final._body._element.append(copy.deepcopy(element))

    doc_final.save(caminho_saida)

# --- Processadores ---

def processar_documentos_individual(df, template_path, pasta_saida, nome_modelo_higienizado, coluna_nomeacao, callback_progresso):
    arquivos_gerados = []
    total = len(df)
    for index, row in df.iterrows():
        doc = Document(template_path)
        row_dados = row.to_dict()

        for p in doc.paragraphs:
            substituir_preservando_formatacao(p, row_dados)

        for table in doc.tables:
            for r in table.rows:
                for cell in r.cells:
                    for p_cell in cell.paragraphs:
                        substituir_preservando_formatacao(p_cell, row_dados)

        identificador_linha = higienizar_nome_arquivo(row_dados.get(coluna_nomeacao, f"linha_{index}"))
        nome_arquivo_saida = os.path.join(
            pasta_saida,
            f"{nome_modelo_higienizado}_{index + 1}_{identificador_linha}.docx"
        )
        doc.save(nome_arquivo_saida)
        arquivos_gerados.append(nome_arquivo_saida)
        
        callback_progresso((index + 1) / total * 50)
        
    return arquivos_gerados

def processar_documentos_agrupados(df, template_path, pasta_saida, nome_modelo_higienizado, coluna_grupo, callback_progresso):
    grupos = df[coluna_grupo].dropna().unique()
    arquivos_gerados = []
    total_grupos = len(grupos)

    for idx_g, grupo in enumerate(grupos):
        df_grupo = df[df[coluna_grupo] == grupo].reset_index(drop=True)
        doc = Document(template_path)

        valor_grupo_str = str(grupo).strip()
        mapeamento_grupo = {coluna_grupo: valor_grupo_str}
        for p in doc.paragraphs:
            substituir_preservando_formatacao(p, mapeamento_grupo)

        paragraphs = list(doc.paragraphs)
        i = 0
        while i < len(paragraphs):
            text = paragraphs[i].text.strip()
            match_start = re.match(r'\{\{iniciar:([^}]+)\}\}(.*)', text)
            if match_start:
                nome_coluna_tag = match_start.group(1).strip()
                start_idx = i
                end_idx = -1

                for j in range(start_idx + 1, len(paragraphs)):
                    text_fim = paragraphs[j].text.strip()
                    if text_fim == f"{{{{fim:{nome_coluna_tag}}}}}":
                        end_idx = j
                        break

                if end_idx != -1:
                    blocos_paragrafos_modelo = paragraphs[start_idx+1:end_idx]

                    for idx_para in range(end_idx, start_idx - 1, -1):
                        p_to_remove = paragraphs[idx_para]
                        p_to_remove._element.getparent().remove(p_to_remove._element)

                    ref_p = paragraphs[start_idx - 1] if start_idx > 0 else None

                    total_registros = len(df_grupo)
                    for idx_registro, row_dados in df_grupo.iterrows():
                        row_dados_dict = row_dados.to_dict()
                        for p_modelo in blocos_paragrafos_modelo:
                            if ref_p:
                                novo_p = duplicar_paragrafo_estilizado(p_modelo, ref_p)
                                ref_p = novo_p
                            else:
                                ref_p = doc.add_paragraph()
                                for run_mod in p_modelo.runs:
                                    run_add = ref_p.add_run(run_mod.text)
                                    if run_mod._r.rPr is not None:
                                        run_add._r.append(OxmlElement('w:rPr'))
                                        for child in run_mod._r.rPr:
                                            run_add._r.rPr.append(child.__copy__())

                            substituir_preservando_formatacao(ref_p, row_dados_dict)

                        if idx_registro < total_registros - 1:
                            if ref_p:
                                novo_p = inserir_paragrafo_depois(ref_p, "")
                                ref_p = novo_p
                            else:
                                ref_p = doc.add_paragraph("")

                    paragraphs = list(doc.paragraphs)
                    i = start_idx - 1
            i += 1

        tabela_alvo = None
        linha_modelo = None

        for table in doc.tables:
            for r_idx, row in enumerate(table.rows):
                for cell in row.cells:
                    if "{{" in cell.text and "}}" in cell.text:
                        tabela_alvo = table
                        linha_modelo = row
                        break
                if tabela_alvo:
                    break

        if tabela_alvo and linha_modelo:
            placeholders_linha = []
            for cell in linha_modelo.cells:
                encontrados = re.findall(r'\{\{([^}]+)\}\}', cell.text)
                placeholders_linha.append(encontrados[0] if encontrados else None)

            tr_elemento = linha_modelo._tr
            tabela_alvo._tbl.remove(tr_elemento)

            for _, row_dados in df_grupo.iterrows():
                nova_linha = tabela_alvo.add_row()
                row_dados_dict = row_dados.to_dict()
                for idx_c, tag_completa in enumerate(placeholders_linha):
                    p_cell = nova_linha.cells[idx_c].paragraphs[0]
                    if tag_completa:
                        substituir_preservando_formatacao(p_cell, row_dados_dict)
                    else:
                        p_cell.text = ""

        elif not any(re.match(r'\{\{iniciar:[^}]+\}\}', p.text.strip()) for p in doc.paragraphs):
            for p in doc.paragraphs:
                row_dados_dict = df_grupo.iloc[0].to_dict()
                substituir_preservando_formatacao(p, row_dados_dict)

        identificador_grupo = higienizar_nome_arquivo(valor_grupo_str)
        nome_arquivo_saida = os.path.join(
            pasta_saida,
            f"{nome_modelo_higienizado}_{identificador_grupo}.docx"
        )
        doc.save(nome_arquivo_saida)
        arquivos_gerados.append(nome_arquivo_saida)
        
        callback_progresso((idx_g + 1) / total_grupos * 50)
        
    return arquivos_gerados

def iniciar_interface():
    root = tk.Tk()
    root.title("GeraDoc v4.2")
    root.geometry("500x520")
    root.resizable(False, False)

    var_excel = tk.StringVar()
    var_modelo = tk.StringVar()
    var_agrupar = tk.BooleanVar(value=False)
    var_arquivo_unico = tk.BooleanVar(value=False)
    var_pdf = tk.BooleanVar(value=False)
    var_coluna_grupo = tk.StringVar()
    var_coluna_nomeacao = tk.StringVar()
    
    var_progresso = tk.DoubleVar()
    var_status = tk.StringVar(value="Aguardando...")

    colunas_disponiveis = []

    def selecionar_excel():
        caminho = filedialog.askopenfilename(
            title="Selecione a Planilha Excel",
            filetypes=[("Arquivos de Excel", "*.xlsx *.xls")]
        )
        if caminho:
            var_excel.set(caminho)
            try:
                df = pd.read_excel(caminho)
                nonlocal colunas_disponiveis
                colunas_disponiveis = [str(col).strip() for col in df.columns]
                
                menu_colunas['values'] = colunas_disponiveis
                menu_nomeacao['values'] = colunas_disponiveis
                
                if colunas_disponiveis:
                    var_coluna_grupo.set(colunas_disponiveis[0])
                    var_coluna_nomeacao.set(colunas_disponiveis[0])
            except Exception as e:
                messagebox.showerror("Erro ao ler Excel", f"Não foi possível carregar as colunas:\n{e}")

    def selecionar_modelo():
        caminho = filedialog.askopenfilename(
            title="Selecione o Modelo Word",
            filetypes=[("Documentos do Word", "*.docx")]
        )
        if caminho:
            var_modelo.set(caminho)

    def atualizar_visibilidade():
        if var_agrupar.get():
            frame_agrupamento.grid()
        else:
            frame_agrupamento.grid_remove()
            
        if var_agrupar.get() or var_arquivo_unico.get():
            frame_nomeacao.grid_remove()
        else:
            frame_nomeacao.grid()
            
    def set_progresso(valor):
        var_progresso.set(valor)
        root.update_idletasks()

    def processamento_em_background():
        excel_path = var_excel.get()
        template_path = var_modelo.get()

        try:
            var_status.set("Lendo a planilha...")
            df = pd.read_excel(excel_path)
            df.columns = df.columns.str.strip()

            output_dir = os.path.dirname(excel_path)
            pasta_saida = os.path.join(output_dir, "Documentos_Gerados")
            os.makedirs(pasta_saida, exist_ok=True)

            nome_modelo_base = os.path.splitext(os.path.basename(template_path))[0]
            nome_modelo_higienizado = higienizar_nome_arquivo(nome_modelo_base)

            arquivos_gerados = []
            arquivos_finais = []

            var_status.set("Gerando arquivos do Word...")
            if var_agrupar.get():
                coluna_escolhida = var_coluna_grupo.get()
                arquivos_gerados = processar_documentos_agrupados(df, template_path, pasta_saida, nome_modelo_higienizado, coluna_escolhida, set_progresso)
            else:
                col_nome = var_coluna_nomeacao.get() if var_coluna_nomeacao.get() else df.columns[0]
                arquivos_gerados = processar_documentos_individual(df, template_path, pasta_saida, nome_modelo_higienizado, col_nome, set_progresso)

            if var_arquivo_unico.get() and len(arquivos_gerados) > 0:
                var_status.set("Juntando tudo em um arquivo único...")
                caminho_unico = os.path.join(pasta_saida, f"{nome_modelo_higienizado}_ARQUIVO_UNICO.docx")
                combinar_documentos(arquivos_gerados, caminho_unico)
                
                for arq in arquivos_gerados:
                    try:
                        os.remove(arq)
                    except Exception:
                        pass
                arquivos_finais = [caminho_unico]
            else:
                arquivos_finais = arquivos_gerados
                
            if not var_pdf.get():
                set_progresso(100)

            aviso_pdf = ""
            if var_pdf.get():
                var_status.set("Convertendo para PDF... (Isso pode demorar)")
                if not HAS_WIN32COM:
                    aviso_pdf = "\n\nAVISO: A conversão para PDF falhou porque a biblioteca 'pywin32' não está instalada."
                    set_progresso(100)
                else:
                    try:
                        total_pdf = len(arquivos_finais)
                        for idx_p, arq in enumerate(arquivos_finais):
                            if os.path.exists(arq):
                                converter_para_pdf_seguro(arq)
                            progresso_atual = 50 + ((idx_p + 1) / total_pdf * 50)
                            set_progresso(progresso_atual)
                            
                        aviso_pdf = "\n\n(Os arquivos também foram convertidos para PDF com sucesso!)"
                    except Exception as e:
                        aviso_pdf = f"\n\nAVISO: Houve um erro ao tentar converter para PDF.\nErro técnico: {e}"
                        set_progresso(100)

            var_status.set("Concluído!")
            msg_final = f"Processamento concluído!\n\nForam gerados {len(arquivos_finais)} arquivo(s) na pasta:\n{pasta_saida}{aviso_pdf}"
            
            root.after(0, lambda: [btn_processar.config(state='normal'), messagebox.showinfo("Sucesso!", msg_final), root.destroy()])

        except Exception as e:
            root.after(0, lambda: [btn_processar.config(state='normal'), var_status.set("Erro no processamento."), messagebox.showerror("Erro no Processamento", f"Ocorreu um erro:\n{e}")])

    def disparar_processamento():
        if not var_excel.get() or not var_modelo.get():
            messagebox.showwarning("Atenção", "Por favor, selecione ambos os arquivos (Excel e Modelo).")
            return
            
        if var_agrupar.get() and not var_coluna_grupo.get():
            messagebox.showwarning("Atenção", "Selecione a coluna para agrupamento.")
            return

        btn_processar.config(state='disabled')
        var_progresso.set(0)
        
        tarefa = threading.Thread(target=processamento_em_background)
        tarefa.start()

    pad_options = {'padx': 10, 'pady': 5}

    lbl_excel = tk.Label(root, text="Planilha de Dados (Excel):")
    lbl_excel.grid(row=0, column=0, sticky='w', **pad_options)
    ent_excel = tk.Entry(root, textvariable=var_excel, width=40, state='readonly')
    ent_excel.grid(row=1, column=0, **pad_options)
    btn_excel = tk.Button(root, text="Procurar", command=selecionar_excel)
    btn_excel.grid(row=1, column=1, **pad_options)

    lbl_modelo = tk.Label(root, text="Documento de Modelo (Word):")
    lbl_modelo.grid(row=2, column=0, sticky='w', **pad_options)
    ent_modelo = tk.Entry(root, textvariable=var_modelo, width=40, state='readonly')
    ent_modelo.grid(row=3, column=0, **pad_options)
    btn_modelo = tk.Button(root, text="Procurar", command=selecionar_modelo)
    btn_modelo.grid(row=3, column=1, **pad_options)

    frame_nomeacao = tk.Frame(root)
    lbl_nomeacao = tk.Label(frame_nomeacao, text="Qual coluna dará nome aos arquivos?")
    lbl_nomeacao.pack(side='left', padx=(0, 5))
    menu_nomeacao = ttk.Combobox(frame_nomeacao, textvariable=var_coluna_nomeacao, state='readonly', width=18)
    menu_nomeacao.pack(side='left')
    frame_nomeacao.grid(row=4, column=0, columnspan=2, sticky='w', pady=(5, 5), padx=10)

    chk_agrupar = tk.Checkbutton(
        root,
        text="Deseja agrupar os dados por alguma coluna?",
        variable=var_agrupar,
        command=atualizar_visibilidade
    )
    chk_agrupar.grid(row=5, column=0, columnspan=2, sticky='w', **pad_options)

    frame_agrupamento = tk.Frame(root)
    lbl_coluna = tk.Label(frame_agrupamento, text="Coluna para agrupamento:")
    lbl_coluna.pack(side='left', padx=(0, 5))
    menu_colunas = ttk.Combobox(frame_agrupamento, textvariable=var_coluna_grupo, state='readonly', width=18)
    menu_colunas.pack(side='left')
    frame_agrupamento.grid(row=6, column=0, columnspan=2, sticky='w', padx=10)
    frame_agrupamento.grid_remove()

    chk_unico = tk.Checkbutton(
        root,
        text="Gerar em arquivo único",
        variable=var_arquivo_unico,
        command=atualizar_visibilidade
    )
    chk_unico.grid(row=7, column=0, columnspan=2, sticky='w', pady=(5, 5), padx=10)

    chk_pdf = pip_pdf = tk.Checkbutton(
        root,
        text="Exportar resultados finais também em PDF",
        variable=var_pdf
    )
    chk_pdf.grid(row=8, column=0, columnspan=2, sticky='w', pady=(5, 5), padx=10)

    frame_progresso = tk.Frame(root)
    lbl_status = tk.Label(frame_progresso, textvariable=var_status, font=("Arial", 9, "italic"), fg="gray")
    lbl_status.pack(side='top', pady=(5, 0))
    barra = ttk.Progressbar(frame_progresso, variable=var_progresso, maximum=100, length=500)
    barra.pack(side='top', pady=(5, 10))
    frame_progresso.grid(row=9, column=0, columnspan=2)

    btn_processar = tk.Button(root, text="Gerar Documentos", command=disparar_processamento, bg="#4CAF50", fg="white", font=("Arial", 10, "bold"))
    btn_processar.grid(row=10, column=0, columnspan=2, pady=10)

    root.mainloop()

if __name__ == '__main__':
    iniciar_interface()